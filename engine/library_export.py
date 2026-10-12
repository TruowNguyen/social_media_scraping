"""Excel export of the same filtered, retained evidence shown in the library."""
from datetime import datetime, timezone, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from engine.content_library import library

REVIEW_LABELS = {"UNREVIEWED": "Chưa đánh giá", "RELEVANT": "Liên quan",
                 "IRRELEVANT": "Không liên quan", "NEEDS_REVIEW": "Cần xem xét"}
VN = timezone(timedelta(hours=7))
HEADERS = ["STT", "Link URL", "Nền tảng", "Group / nguồn", "Nội dung", "Ngày đăng (VN)",
           "Độ chính xác ngày đăng", "Thu thập gần nhất (VN)", "Số lần ghi nhận", "Số lượt crawl",
           "Đánh giá", "Ghi chú", "Cập nhật đánh giá (VN)", "Lượt crawl", "ID bản ghi", "Dự án", "Trạng thái xác định dự án"]


def excel_date(value):
    if not value:
        return None
    try:
        stamp = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return stamp.astimezone(VN).replace(tzinfo=None) if stamp.tzinfo else None
    except (ValueError, AttributeError):
        return None


def export_library(database, filters=None, sources=None, projects=None, conversation_db=None):
    if conversation_db is not None:
        from engine.conversation_library import library as combined_library
        result = combined_library(database, conversation_db, filters, sources, paginate=False, projects=projects)
    else:
        result = library(database, filters, sources, paginate=False, projects=projects)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kho dữ liệu"
    sheet.sheet_view.showGridLines = False
    headers = HEADERS + (['Loại nội dung','UID tác giả','Tên tác giả','ID bài gốc','ID nội dung cha'] if conversation_db is not None else [])
    sheet.append(headers)
    for index, item in enumerate(result["rows"], 1):
        review = item["review"]
        values = [index, item["url"], item["platform"], item["group"], item["text"],
                  excel_date(item["published_at"]), item["publication_precision"],
                  excel_date(item["collected_at"]), item["observation_count"], len(item["run_ids"]),
                  REVIEW_LABELS.get(review["status"], review["status"]), review["note"],
                  excel_date(review.get("updated_at")), "\n".join(item["run_ids"]), item["record_id"],
                  item['project'], item['project_match_status']]
        if conversation_db is not None:
            values += [item.get('node_type','POST'),item.get('author_uid'),item.get('author_name'),item.get('root_post_id'),item.get('parent_node_id')]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(index + 1, column)
            if isinstance(value, str):
                # Persist external text as text, including strings beginning '='.
                # No formulas, external workbook references, or executable hyperlinks.
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
                cell.data_type = "s"
            else:
                cell.value = value
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=column in {4, 5, 12, 14})
            if column in {6, 8, 13}:
                cell.number_format = "dd/mm/yyyy hh:mm:ss"
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="286FA8")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [8, 48, 14, 36, 75, 24, 22, 24, 18, 16, 22, 50, 24, 48, 42, 35, 25]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    if conversation_db is not None:
        for index in range(18,23): sheet.column_dimensions[get_column_letter(index)].width = 28
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
