"""Saved-data workspace: immutable views, context, versioned human review and health.

No acquisition or model calls are made while constructing a view or reading detail.
"""
from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
import hashlib
import json
import sqlite3
import uuid
import threading
import time

from engine.classification_view import query_rows
from engine.conversation_library import library
from src.acquisition.models import utcnow
from src.acquisition.projects import load_projects, project_fields
from engine.registry import Registry
from src.analysis.taxonomy import load_snapshot
from src.storage.conversations import ConversationStore, encode
from engine.workspace_filter_contract import choices as filter_choices,validate as validate_filters


class WorkspaceError(ValueError):
    def __init__(self, code, message, status=400, **details):
        super().__init__(message)
        self.code, self.status, self.details = code, status, details


def review_version(review):
    return hashlib.sha256(encode(review).encode()).hexdigest()


EMPTY_REVIEW = {'status': 'UNREVIEWED', 'note': '', 'updated_at': None}
REVIEW_LABELS = {'UNREVIEWED':'Chưa đánh giá','RELEVANT':'Liên quan','IRRELEVANT':'Không liên quan','NEEDS_REVIEW':'Cần xem xét'}


class Workspace:
    def __init__(self, database, raw_database, config_dir, classifier, service):
        self.database, self.raw_database = Path(database), Path(raw_database)
        self.config_dir, self.classifier, self.service = Path(config_dir), classifier, service
        self.cache={};self.cache_lock=threading.Lock()
        raw = ConversationStore(self.raw_database)
        try:
            raw.db.executescript('''
                CREATE TABLE IF NOT EXISTS dashboard_views(view_id TEXT PRIMARY KEY, fingerprint TEXT NOT NULL,
                    expires_at TEXT NOT NULL, payload TEXT NOT NULL);
                CREATE INDEX IF NOT EXISTS dashboard_views_fingerprint ON dashboard_views(fingerprint);
                CREATE TABLE IF NOT EXISTS dashboard_review_requests(request_id TEXT PRIMARY KEY,payload TEXT NOT NULL);
                CREATE INDEX IF NOT EXISTS node_parent ON content_nodes(parent_node_id);
            ''')
        finally:
            raw.close()

    def _open(self):
        raw = ConversationStore(self.raw_database, initialize=False)
        raw.db.execute('BEGIN')
        return raw

    def catalog(self):
        registry=Registry(self.config_dir);taxonomy=load_snapshot(self.classifier.taxonomy_dir)
        return {'filter_options':filter_choices(),'sources':[{'source_id':sid,'name':s['name'],'platform':s['platform'],'source_type':s['source_type']} for sid,s in registry.sources.items() if s['source_type'] in {'FACEBOOK_GROUP','THREADS_SEARCH'}],
                'projects':[{'project_id':p['project_id'],'name':p['name']} for p in taxonomy['profiles'].values()],
                'issues':[{'issue_id':i['issue_id'],'name':i['name'],'project_ids':[p['project_id'] for p in taxonomy['profiles'].values() if i['issue_id'] in p['issue_priorities']]} for i in taxonomy['catalog']['issues']]}

    def _review(self, raw, node_id):
        row = raw.db.execute('SELECT payload FROM node_reviews WHERE node_id=?', (node_id,)).fetchone()
        review = json.loads(row[0]) if row else dict(EMPTY_REVIEW)
        return {**review, 'version': review_version(review)}

    @staticmethod
    def _page(payload, page=1, page_size=25):
        page, page_size = int(page), int(page_size)
        if page < 1 or not 1 <= page_size <= 100:
            raise ValueError('Invalid pagination')
        page = min(page, max(1, (payload['total'] + page_size - 1)//page_size))
        return {**{k:v for k,v in payload.items() if k!='rows'}, 'page':page, 'page_size':page_size,
                'rows':payload['rows'][(page-1)*page_size:page*page_size]}

    def stored_view(self, view_id):
        raw = self._open()
        try:
            row = raw.db.execute('SELECT payload,expires_at FROM dashboard_views WHERE view_id=?', (view_id,)).fetchone()
            if not row or row[1] <= utcnow():
                raise WorkspaceError('VIEW_EXPIRED','Phiên kết quả đã hết hạn. Làm mới dữ liệu để tạo phiên mới.',410)
            return json.loads(row[0])
        finally:
            raw.close()

    def view(self, params):
        if params.get('view_id'):
            return self._page(self.stored_view(params['view_id']), params.get('page',1), params.get('page_size',25))
        kind = params.get('kind','CLASSIFY')
        if kind not in {'CLASSIFY','LIBRARY'}:
            raise ValueError('Invalid view kind')
        filters = {k:v for k,v in params.items() if k not in {'kind','page','page_size','refresh','view_id'} and v}
        filters.setdefault('date_basis','published_at' if kind=='CLASSIFY' else 'observed_any')
        filters.setdefault('project_basis','MODEL' if kind=='CLASSIFY' else 'SOURCE')
        filters.setdefault('sort','newest')
        cache_key=encode({'kind':kind,'filters':filters})
        with self.cache_lock:
            cached=self.cache.get(cache_key)
            if params.get('refresh')!='1' and cached and time.monotonic()-cached[0]<15:
                return self._page(cached[1],params.get('page',1),params.get('page_size',25))
        if filters.get('date_basis','published_at') not in {'published_at','observed_any','observed_latest','collected_at'}:raise ValueError('Invalid date basis')
        if filters.get('quality','') not in {'','COMPLETE','PARTIAL','FAILED','REFERENCE_ONLY'}:raise ValueError('Invalid quality')
        if filters.get('project_basis','MODEL') not in {'MODEL','SOURCE'}:raise ValueError('Invalid project basis')
        registry, projects = Registry(self.config_dir), load_projects(self.config_dir)
        taxonomy = load_snapshot(self.classifier.taxonomy_dir)
        try:validate_filters(filters,registry,taxonomy)
        except ValueError as error:raise WorkspaceError('INVALID_FILTER',str(error)) from error
        raw = self._open()
        try:
            # Classification, reviews and raw nodes are captured on the same SQLite read transaction.
            from engine.classification_worker import classification_rows
            reviews = {r[0]:json.loads(r[1]) for r in raw.db.execute('SELECT node_id,payload FROM node_reviews')}
            metadata = {n['node_id']:n for n in raw.list_nodes()}
            query_scopes={}
            for record in raw.db.execute('SELECT node_id,payload FROM conversation_observations'):
                occurrence=json.loads(record[1])
                if occurrence.get('project_scope_id') in projects:query_scopes.setdefault(record[0],set()).add(occurrence['project_scope_id'])
            from src.storage.image_evidence import read,summary as ocr_summary
            for node in metadata.values():node.update(ocr_summary(read(raw,node['node_id']),node['media_refs']))
            if kind == 'CLASSIFY':
                rows = classification_rows(raw, taxonomy)
                for row in rows:
                    node = metadata[row['node_id']]
                    row.update(record_id='v2:'+row['node_id'], revision_id=node['revision_id'],
                        parent_node_id=node['parent_node_id'], text=row['raw_text'],
                        **{k:node[k] for k in ('image_evidence','ocr_text','ocr_status','ocr_processed_images','ocr_observed_images')},
                        author_name=node.get('author_name'), author_uid=node.get('author_uid'),
                        extraction_status=node['extraction_status'], missing_data_reasons=node['missing_fields'],
                        groups=[registry.sources.get(s,{}).get('name',s) for s in row['source_ids']],
                        review=reviews.get(row['node_id'],dict(EMPTY_REVIEW)))
                    row['review']={**row['review'],'version':review_version(row['review'])}
                    scoped=set()
                    for sid in row['source_ids'] or [None]:
                        fields=project_fields({'raw_text':node['raw_text'],'context_source_id':sid},projects)
                        scoped.update(fields['project_ids']);scoped.update(fields['project_context_ids'])
                    if node['platform']=='THREADS':scoped.update(query_scopes.get(node['node_id'],set()))
                    row['source_project_ids']=sorted(scoped)
                local = dict(filters)
                if local.get('date_basis')=='observed_latest':local['date_basis']='collected_at'
                if filters.get('project_basis')=='SOURCE':
                    selected = local.pop('project_id','')
                    if selected:
                        rows=[r for r in rows if (not r['source_project_ids'] if selected=='__unknown__' else selected in r['source_project_ids'])]
                if filters.get('review'):
                    if filters['review'] not in REVIEW_LABELS: raise ValueError('Invalid review')
                    rows=[r for r in rows if r['review']['status']==filters['review']]
                if filters.get('quality'):
                    rows=[r for r in rows if r['extraction_status']==filters['quality']]
                if filters.get('ocr_status'):
                    rows=[r for r in rows if r.get('ocr_status')==filters['ocr_status']]
                if filters.get('date_basis')=='observed_any':
                    # A source-observation window is distinct from the latest collection timestamp.
                    matched = self._observations(raw, filters)
                    rows=[dict(r,matched_observed_at=matched[r['node_id']]) for r in rows if r['node_id'] in matched]
                    local.pop('start_date',None); local.pop('end_date',None);local['date_basis']='collected_at'
                query = query_rows(rows,{**local,'page':1,'page_size':100},paginate=False)
                rows, summary = query['rows'],query['summary']
            else:
                query_filters=dict(filters)
                if query_filters.get('date_basis')=='observed_any':query_filters['date_basis']='collected_at'
                if query_filters.get('date_basis')=='observed_latest':query_filters['date_basis']='latest_collected_at'
                query=library(self.database,self.raw_database,query_filters,registry.sources,paginate=False,projects=projects,raw_store=raw)
                rows=query['rows']
                classifications={r['node_id']:r for r in classification_rows(raw,taxonomy)}
                matched=self._observations(raw,filters) if filters.get('date_basis','observed_any') in {'observed_any','collected_at'} else {}
                for row in rows:
                    node=metadata.get(row['record_id'].removeprefix('v2:'))
                    if node:
                        review=reviews.get(node['node_id'],dict(EMPTY_REVIEW))
                        row.update(node_id=node['node_id'],revision_id=node['revision_id'],review={**review,'version':review_version(review)},matched_observed_at=matched.get(node['node_id']))
                        row.update({k:node[k] for k in ('image_evidence','ocr_text','ocr_status','ocr_processed_images','ocr_observed_images')})
                        analysis=classifications.get(node['node_id'],{})
                        row.update({k:analysis.get(k) for k in ('analysis','processing_status','filter_decision','reason_vi','job_id','model_id','error_code')})
                        row['source_ids']=analysis.get('source_ids',[])
                summary={}
            if kind=='LIBRARY' and filters.get('ocr_status'):
                if filters['ocr_status'] not in {'NO_IMAGES','NOT_RUN','AVAILABLE','EMPTY','FAILED','UNAVAILABLE','NEEDS_REVIEW','PARTIAL'}:raise ValueError('Invalid OCR status')
                rows=[r for r in rows if r.get('ocr_status')==filters['ocr_status']]
            sort=filters.get('sort','newest')
            if sort not in {'newest','oldest'}:raise ValueError('Invalid sort')
            basis=filters.get('date_basis','published_at' if kind=='CLASSIFY' else 'observed_any')
            key='published_at' if basis=='published_at' else 'matched_observed_at' if basis=='observed_any' or basis=='collected_at' and kind=='LIBRARY' else 'collected_at'
            rows.sort(key=lambda r:(r.get(key) or '',r.get('record_id','')),reverse=sort=='newest')
            rows=[r for r in rows if r.get(key)]+[r for r in rows if not r.get(key)]
            for row in rows:
                row['source_refs']=[{'source_id':sid,'name':registry.sources.get(sid,{}).get('name',sid),
                    'url':registry.sources.get(sid,{}).get('url'),'platform':registry.sources.get(sid,{}).get('platform')} for sid in row.get('source_ids',[])]
            if filters.get('run_id') or filters.get('query_id'):
                matched=self._observations(raw,{k:v for k,v in filters.items() if k not in {'start_date','end_date'}})
                rows=[r for r in rows if r.get('node_id') in matched]
            payload={'schema_version':'dashboard-view-v1','kind':kind,'filters':filters,'total':len(rows),'summary':summary,
                'filter_options':filter_choices(),
                'processing_counts':dict(Counter(r['processing_status'] for r in rows)),
                'units':dict(Counter(r['node_type'] for r in rows)), 'unknown_time_count':sum(n['publication_precision']!='INSTANT' for n in metadata.values()),
                'rows':rows,'sources':[{'source_id':sid,'name':s['name'],'platform':s['platform'],'source_type':s['source_type']} for sid,s in registry.sources.items() if s['source_type'] in {'FACEBOOK_GROUP','THREADS_SEARCH'}],
                'projects':[{'project_id':p['project_id'],'name':p['name']} for p in taxonomy['profiles'].values()],
                'issues':[{'issue_id':i['issue_id'],'name':i['name'],'project_ids':[p['project_id'] for p in taxonomy['profiles'].values() if i['issue_id'] in p['issue_priorities']]} for i in taxonomy['catalog']['issues']]}
        finally:
            raw.close()
        fingerprint=hashlib.sha256(encode(payload).encode()).hexdigest()
        raw=ConversationStore(self.raw_database,initialize=False)
        try:
            with raw.transaction():
                raw.db.execute('DELETE FROM dashboard_views WHERE expires_at<=?',(utcnow(),))
                saved=raw.db.execute('SELECT payload FROM dashboard_views WHERE fingerprint=? AND expires_at>? LIMIT 1',(fingerprint,utcnow())).fetchone()
                if saved:payload=json.loads(saved[0])
                else:
                    payload.update(view_id='view_'+uuid.uuid4().hex,as_of=utcnow())
                    expiry=(datetime.now(timezone.utc)+timedelta(hours=24)).isoformat().replace('+00:00','Z')
                    raw.db.execute('INSERT INTO dashboard_views VALUES (?,?,?,?)',(payload['view_id'],fingerprint,expiry,encode(payload)))
                    raw.db.execute('DELETE FROM dashboard_views WHERE view_id IN (SELECT view_id FROM dashboard_views ORDER BY rowid DESC LIMIT -1 OFFSET 128)')
        finally:raw.close()
        with self.cache_lock:
            if len(self.cache)>=32:self.cache.clear()
            self.cache[cache_key]=(time.monotonic(),payload)
        return self._page(payload,params.get('page',1),params.get('page_size',25))

    @staticmethod
    def _observations(raw, filters):
        from datetime import date, time
        vn=timezone(timedelta(hours=7))
        start=datetime.combine(date.fromisoformat(filters['start_date']),time(),vn) if filters.get('start_date') else None
        end=datetime.combine(date.fromisoformat(filters['end_date'])+timedelta(days=1),time(),vn) if filters.get('end_date') else None
        if start and end and start>=end:raise ValueError('Invalid dates')
        result={}
        for row in raw.db.execute('SELECT node_id,payload,source_id,run_id FROM conversation_observations'):
            if filters.get('source_id') and row[2]!=filters['source_id']:continue
            if filters.get('query_id') and row[2]!=filters['query_id']:continue
            if filters.get('run_id') and row[3]!=filters['run_id']:continue
            stamp=json.loads(row[1]).get('observed_at')
            if not stamp:continue
            value=datetime.fromisoformat(stamp.replace('Z','+00:00'))
            if (not start or value>=start) and (not end or value<end):result[row[0]]=max(stamp,result.get(row[0],''))
        return result

    def detail(self,node_id):
        raw=self._open()
        try:
            node=raw.get(node_id)
            if not node:raise WorkspaceError('NOT_FOUND','Không tìm thấy nội dung.',404)
            from src.storage.image_evidence import read,summary as ocr_summary
            node.update(ocr_summary(read(raw,node_id),node['media_refs']))
            ancestors=[];missing=[];seen={node_id};parent=node['parent_node_id']
            while parent and len(ancestors)<100:
                if parent in seen:missing.append('PARENT_CYCLE');break
                seen.add(parent);item=raw.get(parent)
                if not item:missing.append('PARENT_NOT_COLLECTED');break
                item.update(ocr_summary(read(raw,parent),item['media_refs']))
                ancestors.insert(0,item);parent=item['parent_node_id']
            if node['root_post_id'] and node['root_post_id'] not in seen:
                root=raw.get(node['root_post_id'])
                if root:
                    root.update(ocr_summary(read(raw,root['node_id']),root['media_refs']));ancestors.insert(0,root)
                else:missing.append('ROOT_NOT_COLLECTED')
            from engine.classification_worker import classification_rows
            taxonomy=load_snapshot(self.classifier.taxonomy_dir)
            analysis=next((r for r in classification_rows(raw,taxonomy) if r['node_id']==node_id),None)
            observation=raw.db.execute('SELECT payload FROM conversation_observations WHERE node_id=? ORDER BY rowid DESC LIMIT 1',(node_id,)).fetchone()
            registry=Registry(self.config_dir)
            source_refs=[{'source_id':sid,'name':registry.sources.get(sid,{}).get('name',sid),
                'url':registry.sources.get(sid,{}).get('url'),'platform':node['platform'].lower()}
                for (sid,) in raw.db.execute('SELECT source_id FROM node_sources WHERE node_id=? ORDER BY source_id',(node_id,))]
            return {'node':node,'ancestors':ancestors,'missing_context':missing,'classification':analysis,'review':self._review(raw,node_id),
                'source_refs':source_refs,
                'taxonomy_issues':[{'issue_id':i['issue_id'],'name':i['name']} for i in taxonomy['catalog']['issues']],
                'filter_options':filter_choices(),
                'last_observation':json.loads(observation[0]) if observation else None,
                'link_precision':'POST_FALLBACK' if node['platform'].lower()=='facebook' and node['node_type']!='POST' and 'comment_id=' not in (node.get('display_url') or '') else 'ITEM_REFERENCE'}
        finally:raw.close()

    def thread(self,root_id,params):
        offset=int(params.get('offset',0));limit=int(params.get('limit',25))
        if offset<0 or not 1<=limit<=100:raise ValueError('Invalid thread pagination')
        raw=self._open()
        try:
            total=raw.db.execute('SELECT count(*) FROM content_nodes WHERE root_post_id=? AND node_id!=?',(root_id,root_id)).fetchone()[0]
            refs=raw.db.execute('SELECT node_id FROM content_nodes WHERE root_post_id=? AND node_id!=? ORDER BY first_seen_at,node_id LIMIT ? OFFSET ?',(root_id,root_id,limit,offset)).fetchall()
            return {'root':raw.get(root_id),'rows':[raw.get(r[0]) for r in refs],'total':total,'offset':offset,'next_offset':offset+limit if offset+limit<total else None,'coverage':'COLLECTED_ONLY'}
        finally:raw.close()

    def save_reviews(self,payload):
        with self.cache_lock:self.cache.clear()
        changes=payload.get('changes')
        if not isinstance(changes,list) or not 1<=len(changes)<=100:raise ValueError('Invalid selection')
        ids=[c.get('node_id') for c in changes]
        if len(set(ids))!=len(ids):raise ValueError('Duplicate selection')
        request_id=payload.get('request_id')
        if not isinstance(request_id,str) or not 1<=len(request_id)<=128:raise ValueError('Request ID required')
        request_hash=hashlib.sha256(encode(changes).encode()).hexdigest()
        raw=ConversationStore(self.raw_database,initialize=False)
        try:
            with raw.transaction():
                saved=raw.db.execute('SELECT payload FROM dashboard_review_requests WHERE request_id=?',(request_id,)).fetchone()
                if saved:
                    cached=json.loads(saved[0])
                    if cached['request_hash']!=request_hash:raise WorkspaceError('REQUEST_CONFLICT','Mã yêu cầu đã dùng cho một thay đổi khác.',409)
                    return cached['result']
                conflicts=[]
                for change in changes:
                    node=raw.get(change['node_id'])
                    if not node:raise WorkspaceError('NOT_FOUND','Nội dung không còn tồn tại.',404)
                    current=self._review(raw,change['node_id'])
                    if change.get('expected_version')!=current['version'] or change.get('revision_id')!=node['revision_id']:
                        conflicts.append({'node_id':node['node_id'],'review':current,'revision_id':node['revision_id']})
                if conflicts:raise WorkspaceError('REVIEW_CONFLICT','Nội dung hoặc đánh giá đã thay đổi. So sánh với bản đã lưu trước khi cập nhật.',409,conflicts=conflicts)
                result=[]
                for change in changes:
                    raw.review(change['node_id'],change['status'],change.get('note',''))
                    result.append({'node_id':change['node_id'],'review':self._review(raw,change['node_id'])})
                response={'saved':result,'count':len(result)}
                raw.db.execute('INSERT INTO dashboard_review_requests VALUES (?,?)',(request_id,encode({'request_hash':request_hash,'result':response})))
                return response
        finally:raw.close()

    def health(self):
        raw=self._open()
        try:
            tables={r[0] for r in raw.db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            job_counts=dict(raw.db.execute('SELECT status,count(*) FROM classification_jobs GROUP BY status')) if 'classification_jobs' in tables else {}
            from engine.classification_worker import classification_rows
            counts=dict(Counter(r['processing_status'] for r in classification_rows(raw,load_snapshot(self.classifier.taxonomy_dir))))
            reports=[json.loads(r[0]) for r in raw.db.execute('SELECT payload FROM dashboard_raw_runs ORDER BY rowid DESC LIMIT 30')] if 'dashboard_raw_runs' in tables else []
            last_seen=raw.db.execute('SELECT max(last_seen_at) FROM content_nodes').fetchone()[0]
            nodes=raw.db.execute('SELECT count(*) FROM content_nodes').fetchone()[0]
            workers=[json.loads(r[0]) for r in raw.db.execute('SELECT payload FROM dashboard_worker_heartbeats')] if 'dashboard_worker_heartbeats' in tables else []
            for worker in workers:
                worker['health']='STOPPED' if worker.get('status')=='STOPPED' else 'ACTIVE' if (datetime.now(timezone.utc)-datetime.fromisoformat(worker['updated_at'].replace('Z','+00:00'))).total_seconds()<45 else 'STALE'
            workers.sort(key=lambda w:w['updated_at'],reverse=True)
            workers.sort(key=lambda w:w['health']!='ACTIVE')
        finally:raw.close()
        registry=Registry(self.config_dir)
        sources=[]
        for sid,s in registry.sources.items():
            if s['source_type']!='FACEBOOK_GROUP':continue
            latest=next(({'run_id':r['crawl_run_id'],'finished_at':r.get('finished_at'),**j} for r in reports for j in r.get('jobs',[]) if j['source_id']==sid),None)
            sources.append({'source_id':sid,'name':s['name'],'registry_ready':registry.readiness(s) is None,'latest_outcome':latest})
        schedule=self.service.schedule_status() if hasattr(self.service,'schedule_status') else {'enabled':False,'health':'UNKNOWN'}
        threads=self.service.threads.catalog() if getattr(self.service,'threads',None) else None
        return {'as_of':utcnow(),'schedule':schedule,'threads':threads,'classification':{'counts':counts,'job_counts':job_counts,'workers':workers,'health':workers[0]['health'] if workers else 'UNKNOWN',**self.classifier.status()},
            'sources':sources,'runs':reports,'last_collected_at':last_seen,'nodes':nodes}

    def preview(self):
        from engine.classification_worker import classification_rows
        raw=self._open()
        try:
            rows=classification_rows(raw,load_snapshot(self.classifier.taxonomy_dir))
            counts=dict(Counter(r['processing_status'] for r in rows))
            return {'scope':'GLOBAL','cycle_limit':10,'counts':counts,
                    'eligible':sum(counts.get(s,0) for s in ('UNCLASSIFIED','STALE','PENDING')),
                    'notice':'Xử lý một lượt hàng chờ toàn hệ thống, tối đa 10 nội dung; không phân loại lại kết quả còn hiệu lực.'}
        finally:raw.close()

    def retry(self,payload):
        with self.cache_lock:self.cache.clear()
        ids=payload.get('node_ids',[])
        if not isinstance(ids,list) or not 1<=len(ids)<=100 or len(set(ids))!=len(ids):raise ValueError('Invalid retry selection')
        from engine.classification_worker import classification_rows
        raw=ConversationStore(self.raw_database,initialize=False)
        try:
            with raw.transaction():
                current={r['node_id']:r for r in classification_rows(raw,load_snapshot(self.classifier.taxonomy_dir))}
                if any(n not in current or current[n]['processing_status']!='FAILED' or not current[n]['job_id'] for n in ids):
                    raise WorkspaceError('RETRY_NOT_ELIGIBLE','Chỉ thử lại lỗi phân loại còn hiệu lực. Làm mới để kiểm tra trạng thái.',409)
                for n in ids:
                    raw.db.execute("UPDATE classification_jobs SET status='PENDING',attempts=0,error_code=NULL,next_attempt_at=NULL,lease_token=NULL,lease_until=NULL WHERE job_id=? AND status='FAILED'",(current[n]['job_id'],))
            return {'queued':len(ids),'notice':'Đã đưa về hàng chờ; worker sẽ xử lý khi chạy.'}
        finally:raw.close()

    def export(self,view_id):
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        payload=self.stored_view(view_id);workbook=Workbook();sheet=workbook.active
        sheet.title='Phản ánh' if payload['kind']=='CLASSIFY' else 'Kho dữ liệu'
        headings=['ID nội dung','Loại','Link URL','Group','UID tác giả','Tên tác giả','Ngày đăng UTC','Thời gian xác minh','Nội dung','Dự án','Tóm tắt vấn đề','Quyết định bộ lọc','Kết quả model','Trạng thái xử lý','Lý do','Đánh giá người','Ghi chú đã lưu','Thời điểm lưu UTC','Chất lượng','Thiếu dữ liệu','ID bài gốc','ID nội dung cha','Phiên bản nội dung']
        sheet.append(headings)
        sheet.cell(1,len(headings)+1,'Chữ trong ảnh');sheet.cell(1,len(headings)+2,'Trạng thái OCR');sheet.cell(1,len(headings)+3,'Ảnh xử lý / quan sát')
        sheet.cell(1,len(headings)+4,'Nhóm vấn đề (taxonomy)');sheet.cell(1,len(headings)+5,'Mã nhóm vấn đề');sheet.cell(1,len(headings)+6,'Ý định nội dung')
        sheet.cell(1,len(headings)+7,'Nền tảng');sheet.cell(1,len(headings)+8,'ID nguồn / truy vấn')
        sheet.cell(1,len(headings)+9,'Nguồn / truy vấn')
        issues={i['issue_id']:i['name'] for i in payload['issues']}
        projects={p['project_id']:p['name'] for p in payload['projects']}
        intent_names={i['id']:i['name'] for i in filter_choices()['intents']}
        evidence_sheet=workbook.create_sheet('Chữ trong ảnh')
        evidence_sheet.append(['ID nội dung','ID bằng chứng','Ảnh tham chiếu','Trạng thái','Chất lượng','Phần','Chữ trong ảnh','Thời điểm UTC'])
        for row in payload['rows']:
            analysis=row.get('analysis') or {};review=row.get('review') or EMPTY_REVIEW
            source_names=', '.join(row.get('groups',[])) or row.get('group','')
            values=[row.get('node_id',row.get('record_id')),row['node_type'],row['url'],'' if row.get('platform')=='threads' else source_names,row.get('author_uid'),row.get('author_name'),row.get('published_at'),row.get('publication_precision'),row.get('text'),
                ', '.join(projects.get(p['project_id'],p['project_id']) for p in analysis.get('project_matches',[])) or row.get('project'),'; '.join(i['summary'] for i in analysis.get('issues',[])),row.get('filter_decision',{}).get('prefilter_decision'),analysis.get('relevance'),row.get('processing_status'),row.get('reason_vi'),REVIEW_LABELS[review['status']],review['note'],review['updated_at'],row.get('extraction_status'),', '.join(row.get('missing_data_reasons',[])),row.get('root_post_id'),row.get('parent_node_id'),row.get('revision_id')]
            translations={'POST':'Bài viết','COMMENT':'Bình luận','REPLY':'Trả lời','KEEP':'Giữ','DROP':'Loại','REVIEW':'Cần xem xét','PASS':'Cho xử lý tiếp','PENDING':'Đang chờ','RUNNING':'Đang xử lý','DONE':'Hoàn thành','FAILED':'Lỗi','STALE':'Cần phân loại lại','MISSING_TEXT':'Thiếu nội dung để phân loại','OVERSIZED':'Vượt kích thước','UNCLASSIFIED':'Chưa phân loại','FILTER_DROP':'Bộ lọc loại','COMPLETE':'Đầy đủ','PARTIAL':'Thiếu dữ liệu','REFERENCE_ONLY':'Chỉ có tham chiếu','INSTANT':'Đã xác minh giờ đăng','UNKNOWN':'Chưa xác minh'}
            for index in (1,7,11,12,13,18):values[index]=translations.get(values[index],values[index])
            sheet.append(values)
            text=row.get('ocr_text') or ''
            sheet.cell(sheet.max_row,len(headings)+1,text if len(text)<=32700 else text[:32600]+'\n[Xem đầy đủ tại sheet Chữ trong ảnh]')
            sheet.cell(sheet.max_row,len(headings)+2,row.get('ocr_status','NOT_RUN'))
            sheet.cell(sheet.max_row,len(headings)+3,f"{row.get('ocr_processed_images',0)} / {row.get('ocr_observed_images',0)}")
            sheet.cell(sheet.max_row,len(headings)+4,'; '.join(issues.get(i['issue_id'],i['issue_id']) for i in analysis.get('issues',[])))
            sheet.cell(sheet.max_row,len(headings)+5,'; '.join(i['issue_id'] for i in analysis.get('issues',[])))
            intents=analysis.get('content_intents') or (row.get('filter_decision',{}).get('content_intents') if row.get('processing_status')=='FILTER_DROP' else []) or []
            sheet.cell(sheet.max_row,len(headings)+6,', '.join(intent_names.get(i,i) for i in intents))
            sheet.cell(sheet.max_row,len(headings)+7,'Threads' if row.get('platform')=='threads' else 'Facebook')
            sheet.cell(sheet.max_row,len(headings)+8,', '.join(row.get('source_ids',[])))
            sheet.cell(sheet.max_row,len(headings)+9,source_names)
            for evidence in row.get('image_evidence',[]):
                text=evidence.get('text') or '';chunks=[text[i:i+32000] for i in range(0,len(text),32000)] or ['']
                for part,chunk in enumerate(chunks,1):evidence_sheet.append([row.get('node_id'),evidence['evidence_id'],evidence.get('media_ref',{}).get('url'),evidence['status'],evidence['quality_status'],part,chunk,evidence.get('processed_at')])
        for line in evidence_sheet:
            for cell in line:
                if isinstance(cell.value,str):cell.value=ILLEGAL_CHARACTERS_RE.sub('',cell.value);cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        evidence_sheet.freeze_panes='A2';evidence_sheet.auto_filter.ref=evidence_sheet.dimensions;evidence_sheet.column_dimensions['G'].width=80
        for line in sheet:
            for cell in line:
                if isinstance(cell.value,str):cell.value=ILLEGAL_CHARACTERS_RE.sub('',cell.value);cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='21679F')
        for i in range(1,len(headings)+1):sheet.column_dimensions[get_column_letter(i)].width=28
        for i in range(len(headings)+1,len(headings)+10):sheet.column_dimensions[get_column_letter(i)].width=32
        sheet.column_dimensions['I'].width=70;sheet.freeze_panes='D2';sheet.auto_filter.ref=sheet.dimensions
        meta=workbook.create_sheet('Phạm vi')
        for key,value in [('view_id',payload['view_id']),('as_of',payload['as_of']),('kind',payload['kind']),('count',payload['total']),('filters',encode(payload['filters'])),('ghi_chu','Chỉ đánh giá/ghi chú đã lưu tại thời điểm tạo phiên; gồm mọi trang; timestamp UTC.')]:
            meta.append([key,value])
        for line in meta:
            for cell in line:
                if isinstance(cell.value,str):cell.data_type='s'
        meta.column_dimensions['A'].width=22;meta.column_dimensions['B'].width=110
        output=BytesIO();workbook.save(output);return output.getvalue()
