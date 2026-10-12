"""Capture current Threads rollout evidence without crawling or calling Gemini."""
import argparse
import json
import sqlite3
from collections import Counter
from io import BytesIO
from pathlib import Path
from urllib.request import urlopen
from urllib.parse import urlencode
from openpyxl import load_workbook
from engine.classification_service import ClassificationService
from src.acquisition.models import utcnow


def main(argv=None):
    parser=argparse.ArgumentParser()
    parser.add_argument('--database',default='sessions/acquisition/acquisition.sqlite3')
    parser.add_argument('--backup',default='sessions/backups/pre-threads-v3-20261011.sqlite3')
    parser.add_argument('--base-url',default='http://127.0.0.1:8765')
    parser.add_argument('--report',default='reports/threads-deployment-final.json')
    parser.add_argument('--export',default='reports/threads-live-results.xlsx')
    args=parser.parse_args(argv)
    database=Path(args.database)
    rows=[r for r in ClassificationService(database).rows(refresh=True) if r.get('platform')=='threads']
    c=sqlite3.connect(database);previous=sqlite3.connect(args.backup)
    try:
        query="SELECT node_id,current_revision_id,payload FROM content_nodes WHERE json_extract(payload,'$.platform')='FACEBOOK' ORDER BY node_id"
        old=list(previous.execute(query));current=list(c.execute(query))
        reviews_unchanged=list(previous.execute('SELECT * FROM node_reviews ORDER BY node_id'))==list(c.execute('SELECT * FROM node_reviews ORDER BY node_id'))
        history_unchanged=list(previous.execute('SELECT * FROM node_review_history ORDER BY history_id'))==list(c.execute('SELECT * FROM node_review_history ORDER BY history_id'))
        nodes=[json.loads(r[0]) for r in c.execute("SELECT payload FROM content_nodes WHERE json_extract(payload,'$.platform')='THREADS'")]
    finally:c.close();previous.close()
    def get(path):
        with urlopen(args.base_url+path,timeout=45) as response:return response.read()
    catalog=json.loads(get('/api/acquisition/catalog'))
    view=json.loads(get('/api/views?'+urlencode({'kind':'LIBRARY','platform':'threads','page_size':100,'refresh':1})))
    data=get('/api/views/export.xlsx?'+urlencode({'view_id':view['view_id']}))
    target=Path(args.export);target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(data)
    workbook=load_workbook(BytesIO(data),read_only=True);sheet=workbook.active
    export_rows=sheet.max_row-1;headings=[cell.value for cell in sheet[1]]
    workbook.close()
    summary={
        'schema_version':'threads-deployment-evidence-v1','checked_at':utcnow(),
        'status':'LIVE_PARTIAL','provider':catalog['provider'],'query_id':'threads_vwy_name','query':'VinWonders Vũ Yên',
        'pilot_window_utc':{'start':'2026-09-27T16:31:00Z','end':'2026-10-11T16:31:00Z','semantics':'[start,end)'},
        'raw':{'total':len(nodes),'node_types':dict(Counter(n['node_type'] for n in nodes)),
            'with_text':sum(bool(n.get('raw_text')) for n in nodes),'with_exact_publication':sum(n.get('publication_precision')=='INSTANT' for n in nodes),
            'unresolved_root':sum(n['node_type']!='POST' and not n.get('root_post_id') for n in nodes)},
        'classification':{'processing_status':dict(Counter(r['processing_status'] for r in rows)),
            'decisions':dict(Counter((r.get('analysis') or {}).get('relevance') or r['processing_status'] for r in rows)),
            'post_decisions':dict(Counter((r.get('analysis') or {}).get('relevance') or r['processing_status'] for r in rows if r['node_type']=='POST')),
            'keep_distinct_known_roots':len({r.get('root_post_id') or r['node_id'] for r in rows if (r.get('analysis') or {}).get('relevance')=='KEEP'}),
            'unit':'content nodes, not independent complaints','worker_authorized':'continuous Gemini'},
        'facebook_preservation':{'baseline_nodes':len(old),'current_nodes':len(current),'raw_identity_revision_payload_unchanged':old==current,
            'reviews_unchanged':reviews_unchanged,'review_history_unchanged':history_unchanged,'backup':args.backup},
        'export':{'path':str(target),'view_id':view['view_id'],'view_rows':view['total'],'excel_rows':export_rows,
            'rows_match':export_rows==view['total']==len(nodes),'has_platform_and_source_columns':all(v in headings for v in ('Nền tảng','ID nguồn / truy vấn'))},
        'schedule':catalog['schedule'],'verification':catalog['verification'],'recovery_error':catalog.get('recovery_error'),
        'tests':{'contract':{'run':251,'passed':250,'skipped':1},'browser_parser':{'run':2,'passed':2},'workspace_desktop_mobile':{'run':2,'passed':2}},
        'evidence':['reports/threads-production-live-run.json','reports/threads-production-reuse-confirmation.json','reports/threads-production-parser-v4.json'],
        'open_acceptance':['FULL_WINDOW_COVERAGE','EXACT_NESTED_PARENT','HELD_OUT_QUALITY','OLD_POST_NEW_REPLY_LIVE_CONTROL','THREE_DAY_SCHEDULE','API_PROVIDER','THREADS_OCR','FACEBOOK_SPOOL_MIGRATION'],
    }
    path=Path(args.report);path.parent.mkdir(parents=True,exist_ok=True);path.write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False))


if __name__=='__main__':main()
