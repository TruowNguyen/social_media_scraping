from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import shutil
import yaml
from unittest.mock import patch
from openpyxl import load_workbook
from src.acquisition.projects import load_projects, project_fields
from src.browser.profiles import PROJECT_ROOT
from src.storage.acquisition import AcquisitionStore
from src.acquisition.models import Candidate, Window, parse_utc
from engine.content_library import library
from engine.library_export import export_library
from engine.dashboard import SearchService, search_output


class ProjectTests(unittest.TestCase):
    def setUp(self):
        self.projects = load_projects(PROJECT_ROOT / 'config')
        self.pid = 'vinwonders_vu_yen'

    def test_aliases_and_boundaries_require_actual_item_evidence(self):
        for caption in ['VinWonder Vũ Yên nhân viên không phản hồi', 'vinwonders vu yen đông quá', '#VinWondersVuYen xếp hàng lâu']:
            self.assertEqual(project_fields({'caption': caption}, self.projects)['project_ids'], [self.pid])
        for payload in [
            {'caption': 'Vũ Yên kẹt xe'}, {'caption': 'Royal Island thu phí'},
            {'caption': 'VinWonders Nha Trang đông'}, {'caption': 'abcVinWonder Vũ Yênxyz'},
            {'search_keyword': 'VinWonder Vũ Yên', 'search_snippet': 'VinWonder Vũ Yên'}, {},
        ]:
            self.assertEqual(project_fields(payload, self.projects)['project_ids'], [])
        self.assertEqual(project_fields({'hashtags': ['#VinWondersVuYen']}, self.projects)['project_ids'], [self.pid])

    def test_shared_alias_does_not_auto_assign_both_projects(self):
        projects = deepcopy(self.projects)
        projects['other'] = dict(projects[self.pid], project_id='other', name='Other')
        self.assertEqual(project_fields({'caption': 'VinWonders Vũ Yên'}, projects)['project_ids'], [])

    def test_library_search_export_filter_without_duplicate_or_storage_mutation(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'db.sqlite3'
            store = AcquisitionStore(path)
            for identifier, caption in [('1', 'VinWonder Vũ Yên nhân viên không phản hồi'), ('2', 'Vũ Yên đang kẹt xe'), ('3', 'VinWonders Nha Trang')]:
                url = 'https://www.facebook.com/groups/123/posts/' + identifier
                store.save_candidate(Candidate('facebook', url, url, 'group', 'run', 'BROWSER',
                    platform_item_id=identifier, caption=caption, published_at='2026-10-09T01:00:00Z', publication_precision='INSTANT'))
            original = store.db.execute('SELECT payload FROM records').fetchall()
            sources = {'group': dict(source_type='FACEBOOK_GROUP', url='https://www.facebook.com/groups/123', name='VinWonders Vũ Yên')}
            result = library(path, {'project_id': self.pid}, sources, projects=self.projects)
            self.assertEqual(result['total'], 1)
            self.assertEqual(result['rows'][0]['project'], 'VinWonders Vũ Yên')
            self.assertEqual(library(path, {'project_id': '__unknown__'}, projects=self.projects)['total'], 2)
            sheet = load_workbook(BytesIO(export_library(path, {'project_id': self.pid}, sources, self.projects))).active
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(sheet['P2'].value, 'VinWonders Vũ Yên')
            self.assertEqual(sheet['Q2'].value, 'MATCHED')
            window = Window(parse_utc('2026-10-08T17:00:00Z'), parse_utc('2026-10-09T10:00:00Z'))
            rows = search_output(path, window=window, projects=self.projects, project_id=self.pid)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]['project'], 'VinWonders Vũ Yên')
            self.assertEqual(store.db.execute('SELECT payload FROM records').fetchall(), original)
            with self.assertRaises(ValueError): library(path, {'project_id': 'missing'}, projects=self.projects)
            store.close()

    def test_search_query_uses_project_and_crawl_rejects_unmapped_groups(self):
        class InlineThread:
            def __init__(self, target, **kwargs): self.target = target
            def start(self): self.target()
        with tempfile.TemporaryDirectory() as directory:
            service = SearchService(Path(directory) / 'db.sqlite3', PROJECT_ROOT / 'config')
            with patch('engine.dashboard.threading.Thread', InlineThread), patch('engine.dashboard.AcquisitionEngine') as engine, patch('engine.dashboard_conversations.crawl') as raw_crawl, patch('engine.dashboard_conversations.results',return_value=[]):
                raw_crawl.return_value = dict(status='PARTIAL',crawl_run_id='v2',jobs=[],unresolved_time=0)
                engine.return_value.run.return_value = dict(status='SUCCESS', crawl_run_id='run', jobs=[])
                job = service.start({'project_id': self.pid, 'keywords': ['không phản hồi']})
                self.assertEqual(job['project_id'], self.pid)
                registry = engine.call_args.args[0]
                self.assertEqual(registry.queries['vincons_ecosystem']['discovery_terms'], ['VinWonders Vũ Yên không phản hồi'])
                service.start({'project_id': self.pid}, 'CRAWL')
                self.assertEqual(set(raw_crawl.call_args.args[2]), set(self.projects[self.pid]['source_ids']))
                self.assertEqual(len(raw_crawl.call_args.args[2]), 7)
            schedule = service.configure_schedule(dict(enabled=True, project_id=self.pid))
            self.assertEqual(schedule['project_id'], self.pid)
            service.configure_schedule(dict(enabled=False))
            config = Path(directory) / 'config'
            shutil.copytree(PROJECT_ROOT / 'config', config)
            data = yaml.safe_load((config / 'projects.yaml').read_text(encoding='utf8'))
            next(p for p in data['projects'] if p['project_id']==self.pid)['source_ids'] = []
            (config / 'projects.yaml').write_text(yaml.safe_dump(data, allow_unicode=True), encoding='utf8')
            service.config_dir = config
            with self.assertRaisesRegex(ValueError, 'group Facebook'): service.start({'project_id': self.pid}, 'CRAWL')
            with self.assertRaises(ValueError):
                service.configure_schedule(dict(enabled=True, project_id=self.pid))
            self.assertFalse(service.schedule['enabled'])
            service.close()

    def test_group_context_remains_unconfirmed_and_explicit_other_project_wins(self):
        from src.acquisition.projects import selected_match
        projects = deepcopy(self.projects)
        projects[self.pid]['source_ids'] = ['mygroup']
        fields = project_fields({'source_id': 'mygroup', 'caption': 'Chờ cả tiếng chưa đến lượt'}, projects)
        self.assertEqual(fields['project_match_status'], 'CONTEXT_ONLY')
        self.assertEqual(fields['project_ids'], [])
        self.assertTrue(selected_match(fields, self.pid))
        projects['other'] = dict(projects[self.pid], project_id='other', name='Other', aliases=['Dự án B'], hashtags=[], source_ids=[])
        fields = project_fields({'source_id': 'mygroup', 'caption': 'Dự án B phục vụ chậm'}, projects)
        self.assertFalse(selected_match(fields, self.pid))

    def test_ticket_sale_exclusion_preserves_customer_refund_complaints(self):
        from src.acquisition.prefilter import classify_payload
        for caption in ['Bán vé VinWonders Vũ Yên, inbox', 'Nhận đặt vé VinWonder Vũ Yên']:
            self.assertEqual(classify_payload({'caption': caption}).prefilter_decision, 'DROP')
        for caption in ['Đại lý vé không hoàn tiền cho tôi', 'Mua vé VinWonders Vũ Yên nhưng không nhận được vé', 'Bán vé giá rẻ nhưng lừa đảo, cảnh báo mọi người']:
            self.assertEqual(classify_payload({'caption': caption}).prefilter_decision, 'REVIEW')
