import copy
import json
import shutil
import tempfile
import threading
import unittest
from datetime import datetime,timedelta,timezone
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace
from io import BytesIO
from openpyxl import load_workbook
from engine.threads_registry import ThreadsRegistry
from engine.threads_worker import crawl
from engine.threads_service import ThreadsService,ThreadsError,ticks
from engine.classification_input import build_item,make_request,build_batches
from engine.classification_service import ClassificationService
from engine.dashboard_workspace import Workspace
from src.analysis.taxonomy import load_snapshot
from src.domain.conversation import ContentNode
from src.acquisition.models import Window,Deadline,AdapterError,utcnow
from src.acquisition.identity import identify
from src.acquisition.threads import node_from_item,node_key,ThreadsAPIAdapter
from src.storage.conversations import ConversationStore
from src.storage.ingest import IngestQueue

ROOT=Path(__file__).resolve().parents[2]
QUERY='threads_vwy_name'

def item(code='abc',text='VinWonders Vũ Yên chờ 40 phút ở cổng',stamp='2026-10-01T01:00:00Z'):
    return dict(url='https://www.threads.com/@tester/post/'+code,shortcode=code,username='tester',text=text,timestamp=stamp)

class FakeAdapter:
    def __init__(self):self.fetched=[]
    def discover(self,request):yield {'items':[item()],'exhausted':False}
    def fetch_item(self,candidate,context):
        self.fetched.append(candidate['url']);return {'root':candidate,'replies':[item('reply','Chờ 50 phút vẫn chưa vào được cổng.')], 'error_code':'DOM_PARENT_MAPPING_UNVERIFIED'}

class ThreadsTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name);self.db=self.path/'raw.sqlite3'
        self.config=self.path/'config';shutil.copytree(ROOT/'config',self.config)
        registry=ThreadsRegistry(self.config);self.data=copy.deepcopy(registry.data)
        self.data.update(enabled=True,verification=dict(search_recent='SUPPORTED',public_replies='SUPPORTED'))
        self.data['limits'].update(seconds_per_run=20,seconds_per_query=20)
        registry.write(self.data,registry.data['version'])
        self.query=next(q for q in self.data['queries'] if q['query_id']==QUERY)
    def tearDown(self):self.tmp.cleanup()
    def raw(self):return ConversationStore(self.db)

    def test_shortcode_is_not_native_id_and_verified_hosts_only(self):
        target=identify(item()['url']);self.assertIsNone(target.item_id)
        self.assertEqual(target.platform,'threads');self.assertEqual(target.kind,'CONTENT_ITEM')
        self.assertNotEqual(node_key(item()),node_key({**item(),'id':'abc'}))
        with self.assertRaises(ValueError):identify('https://threads.com.evil.invalid/@x/post/abc')
    def test_mixed_store_preserves_facebook_and_distinct_platform_identity(self):
        raw=self.raw()
        try:
            fb=raw.ingest(ContentNode('POST',platform_item_id='123',raw_text='caption'), 'fb_source','fb')['node_id']
            before=raw.get(fb);raw.review(fb,'RELEVANT','giữ ghi chú')
            th=raw.ingest(node_from_item({**item(),'id':'123'}),QUERY,'th')['node_id']
            self.assertNotEqual(fb,th);self.assertEqual(raw.get(fb),before)
            self.assertEqual(raw.db.execute("SELECT value FROM conversation_meta WHERE key='schema_version'").fetchone()[0],'conversation-store-v3')
            self.assertEqual(raw.db.execute('SELECT count(*) FROM node_review_history').fetchone()[0],1)
        finally:raw.close()
    def test_first_mixed_write_rolls_back_schema_upgrade(self):
        raw=self.raw()
        try:
            with self.assertRaises(RuntimeError):
                with raw.transaction():raw.ingest(node_from_item(item()),QUERY,'run');raise RuntimeError()
            self.assertEqual(raw.db.execute("SELECT value FROM conversation_meta WHERE key='schema_version'").fetchone()[0],'conversation-store-v2')
            self.assertEqual(len(raw.list_nodes()),0)
        finally:raw.close()
    def test_search_cannot_demote_known_reply_and_type_can_be_refined(self):
        raw=self.raw()
        try:
            root=node_from_item(item(),relationship='UNRESOLVED');raw.ingest(root,QUERY,'r')
            row=node_from_item(item('child'),relationship='UNRESOLVED');raw.ingest(row,QUERY,'r')
            child=node_from_item(item('child'),root=root.node_id,node_type='REPLY',relationship='UNRESOLVED');raw.ingest(child,QUERY,'r')
            raw.ingest(row,QUERY,'r2');saved=raw.get(child.node_id)
            self.assertEqual(saved['node_type'],'REPLY');self.assertEqual(saved['root_post_id'],root.node_id)
        finally:raw.close()
    def test_signed_media_url_refresh_is_not_a_content_revision(self):
        raw=self.raw()
        try:
            a={**item(),'media_refs':[{'kind':'IMAGE','url':'https://instagram.edge1.fna.fbcdn.net/v/photo_123.jpg?signature=old'}]}
            b={**a,'media_refs':[{'kind':'IMAGE','url':'https://instagram.edge2.fna.fbcdn.net/v/photo_123.jpg?signature=new'}]}
            first=raw.ingest(node_from_item(a),QUERY,'r1');second=raw.ingest(node_from_item(b),QUERY,'r2')
            self.assertFalse(second['changed']);self.assertEqual(first['revision_id'],second['revision_id'])
            self.assertEqual(raw.get(first['node_id'])['media_refs'],b['media_refs'])
            b['media_refs'][0]['alt_text']='Description'
            richer=raw.ingest(node_from_item(b),QUERY,'r3')
            b['media_refs'][0]['alt_text']=None
            self.assertFalse(raw.ingest(node_from_item(b),QUERY,'r4')['changed'])
            self.assertEqual(raw.get(first['node_id'])['media_refs'][0]['alt_text'],'Description')
            b['text']='Nội dung đã sửa';self.assertTrue(raw.ingest(node_from_item(b),QUERY,'r3')['changed'])
        finally:raw.close()
    def entries(self):
        return [dict(node=node_from_item(item()).payload(),source_id=QUERY,run_id='run',observed_at=utcnow(),lineage={'query_id':QUERY})]
    def test_page_commit_crash_after_ack_replays_without_observation_or_revision(self):
        queue=IngestQueue(self.path/'spool');path=queue.enqueue(self.entries(),page_id='one',task_id='task',checkpoint={'cursor':'a'})
        raw=self.raw()
        try:
            with patch('src.storage.ingest.os.replace',side_effect=OSError('crash after commit')):
                with self.assertRaises(OSError):queue.commit(path,raw)
            queue.commit(path,raw)
            self.assertEqual(raw.db.execute('SELECT count(*) FROM conversation_observations').fetchone()[0],1)
            self.assertEqual(raw.db.execute('SELECT count(*) FROM content_revisions').fetchone()[0],1)
            self.assertEqual(json.loads(raw.db.execute('SELECT payload FROM conversation_checkpoints').fetchone()[0]),{'cursor':'a'})
        finally:raw.close()
    def test_checksum_failure_does_not_ack_or_commit(self):
        queue=IngestQueue(self.path/'spool');path=queue.enqueue(self.entries(),page_id='one',task_id='task',checkpoint={})
        value=json.loads(path.read_text(encoding='utf8'));value['body']['entries'][0]['node']['raw_text']='tampered';path.write_text(json.dumps(value),encoding='utf8')
        with self.assertRaises(ValueError):queue.drain(self.db)
        raw=self.raw()
        try:self.assertEqual(len(raw.list_nodes()),0);self.assertEqual(raw.db.execute('SELECT count(*) FROM ingest_page_receipts').fetchone()[0],0)
        finally:raw.close()
    def test_expired_lease_cannot_commit_raw_or_checkpoint(self):
        service=ThreadsService(self.db,self.config);raw=self.raw()
        try:
            with raw.transaction():raw.db.execute('INSERT INTO acquisition_leases VALUES (?,?,?,?)',('THREADS','new','newrun',utcnow()))
            queue=IngestQueue(self.path/'spool');path=queue.enqueue(self.entries(),page_id='one',task_id='task',checkpoint={},lease={'platform':'THREADS','token':'old'})
            with self.assertRaisesRegex(RuntimeError,'LEASE_LOST'):queue.commit(path,raw)
            self.assertEqual(len(raw.list_nodes()),0)
        finally:raw.close();service.close()
    def test_raw_before_filter_child_and_rerun_identity(self):
        adapter=FakeAdapter();window=Window(datetime(2026,10,1,tzinfo=timezone.utc),datetime(2026,10,2,tzinfo=timezone.utc))
        first=crawl(self.db,self.data,[self.query],window,'run1',adapter=adapter)
        second=crawl(self.db,self.data,[self.query],window,'run2',adapter=adapter)
        self.assertEqual(first['counters']['committed_nodes'],2);self.assertEqual(second['counters']['unique_new_post'],0)
        raw=self.raw()
        try:self.assertEqual(len(raw.list_nodes()),2);self.assertEqual(raw.db.execute('SELECT count(*) FROM raw_filter_decisions').fetchone()[0],0)
        finally:raw.close()
    def test_crash_recovery_accepts_only_pre_expiry_page(self):
        service=ThreadsService(self.db,self.config);service.close()
        cutoff=datetime.now(timezone.utc)-timedelta(seconds=10)
        raw=self.raw()
        try:
            state={'run_id':'run','status':'COLLECTING','started_at':(cutoff-timedelta(seconds=20)).isoformat()}
            with raw.transaction():
                raw.db.execute('INSERT INTO acquisition_runs VALUES (?,?,?,?,?)',('run','crash','COLLECTING',json.dumps(state),utcnow()))
                raw.db.execute('INSERT INTO acquisition_leases VALUES (?,?,?,?)',('THREADS','dead','run',cutoff.isoformat()))
        finally:raw.close()
        queue=IngestQueue(self.db.parent/'ingest'/self.db.stem)
        entries=self.entries();entries[0]['observed_at']=(cutoff-timedelta(seconds=1)).isoformat()
        queue.enqueue(entries,page_id='before',task_id='crash',checkpoint={'cursor':1},lease={'platform':'THREADS','token':'dead'})
        late=self.entries();late[0]['observed_at']=(cutoff+timedelta(seconds=1)).isoformat();late[0]['node']=node_from_item(item('late')).payload()
        late_path=queue.enqueue(late,page_id='after',task_id='crash',checkpoint={'cursor':2},lease={'platform':'THREADS','token':'dead'})
        recovered=ThreadsService(self.db,self.config)
        raw=self.raw()
        try:
            self.assertEqual(len(raw.list_nodes()),1)
            self.assertEqual(raw.db.execute('SELECT count(*) FROM ingest_page_receipts').fetchone()[0],1)
            self.assertTrue(late_path.exists());self.assertEqual(recovered.catalog()['recovery_error'],'UNCOMMITTED_QUEUE_PAGE')
            self.assertTrue(recovered.get('run')['recovered_after_crash'])
        finally:raw.close();recovered.close()
    def test_recovery_does_not_steal_active_lease(self):
        service=ThreadsService(self.db,self.config);service.close()
        raw=self.raw()
        try:
            with raw.transaction():raw.db.execute('INSERT INTO acquisition_leases VALUES (?,?,?,?)',('THREADS','active','run',(datetime.now(timezone.utc)+timedelta(minutes=1)).isoformat()))
        finally:raw.close()
        queue=IngestQueue(self.db.parent/'ingest'/self.db.stem)
        page=queue.enqueue(self.entries(),page_id='active',task_id='task',checkpoint={},lease={'platform':'THREADS','token':'active'})
        recovered=ThreadsService(self.db,self.config);raw=self.raw()
        try:self.assertTrue(page.exists());self.assertEqual(len(raw.list_nodes()),0)
        finally:raw.close();recovered.close()
    def test_reply_cannot_be_a_conversation_root(self):
        raw=self.raw()
        try:
            root=node_from_item(item(),relationship='UNRESOLVED');raw.ingest(root,QUERY,'r')
            child=node_from_item(item('child'),root=root.node_id,node_type='REPLY',relationship='UNRESOLVED');raw.ingest(child,QUERY,'r')
            grandchild=node_from_item(item('bad'),root=child.node_id,node_type='REPLY',relationship='UNRESOLVED')
            with self.assertRaisesRegex(ValueError,'root'):raw.ingest(grandchild,QUERY,'r')
        finally:raw.close()
    def test_inventory_has_no_fake_window(self):
        result=crawl(self.db,self.data,[self.query],None,'inventory','POST_ONLY',adapter=FakeAdapter())
        self.assertIsNone(result['window_start']);self.assertIsNone(result['window_end']);self.assertEqual(result['task_kind'],'INVENTORY_REFRESH');self.assertNotIn('in_window',result['counters'])
    def test_query_context_is_hint_and_request_v3_batches(self):
        raw=self.raw();snapshot=load_snapshot(ROOT/'specs/taxonomy/v1')
        try:
            n=node_from_item(item(text='Tôi chờ ở cổng'),relationship='UNRESOLVED')
            raw.ingest(n,QUERY,'r',lineage={'query_id':QUERY,'query_version':1,'project_scope_id':'vinwonders_vu_yen'})
            value=build_item(raw,n.node_id,QUERY,snapshot);self.assertEqual(value['project_candidates'][0]['basis'],'QUERY_CONTEXT')
            request=make_request([value],snapshot);self.assertEqual(request['schema_version'],'classification-request-v3')
            self.assertEqual(len(build_batches([value],snapshot)['batches']),1)
        finally:raw.close()
    def test_query_edit_version_compare_and_set(self):
        a,b=ThreadsRegistry(self.config),ThreadsRegistry(self.config)
        change={k:v for k,v in self.query.items() if k!='version'};change['value']='Tên mới'
        a.update_queries([change],a.data['version'])
        with self.assertRaisesRegex(ValueError,'VERSION_CONFLICT'):b.update_queries([change],b.data['version'])
    def test_ticks_include_exact_closing_tick(self):
        self.assertEqual(ticks([{'start':'05:00','end':'08:30'},{'start':'17:00','end':'18:30'}],45),[['05:00','05:45','06:30','07:15','08:00','08:30'],['17:00','17:45','18:30']])
    def test_schedule_legacy_migration_validation_and_cas(self):
        service=ThreadsService(self.db,self.config)
        try:
            old=service.schedule_status();self.assertEqual(old['project_ids'],['vinwonders_vu_yen'])
            for ids in ([],['vincons','vincons'],['missing'],[['invalid']]):
                with self.assertRaises(ThreadsError):service.configure_schedule({'project_ids':ids,'expected_version':old['version']})
            saved=service.configure_schedule({'project_ids':['vincons','vinwonders_vu_yen'],'expected_version':old['version']})
            self.assertFalse(saved['enabled']);self.assertEqual(saved['schema_version'],'platform-schedule-v2')
            with self.assertRaises(ThreadsError):service.configure_schedule({'enabled':True,'expected_version':old['version']})
            single=service.configure_schedule({'project_id':'vincons','expected_version':saved['version']})
            self.assertEqual(single['project_ids'],['vincons'])
        finally:service.close()
    def test_multi_project_ticks_are_serial_idempotent_and_same_window(self):
        import time
        entered=threading.Event();release=threading.Event();calls=[]
        def executor(database,config,queries,window,run_id,scope,**kwargs):
            calls.append((queries[0]['project_id'],window));entered.set();release.wait(5)
            return {'status':'PARTIAL','counters':{}}
        service=ThreadsService(self.db,self.config,executor=executor)
        try:
            s=service.schedule_status();service.configure_schedule({'enabled':True,'project_ids':['vincons','vinwonders_vu_yen'],'expected_version':s['version']})
            moment=datetime(2026,10,1,22,45,tzinfo=timezone.utc)
            first=service.tick(moment);self.assertTrue(entered.wait(2));self.assertIsNone(service.tick(moment))
            self.assertEqual(first['project_id'],'vincons');release.set()
            for _ in range(200):
                if not service.cancels:break
                time.sleep(.01)
            second=service.tick(moment);self.assertEqual(second['project_id'],'vinwonders_vu_yen')
            for _ in range(200):
                if not service.cancels:break
                time.sleep(.01)
            self.assertIsNone(service.tick(moment));self.assertEqual(len(calls),2)
            self.assertEqual(calls[0][1],calls[1][1]);self.assertEqual(first['window_start'],second['window_start'])
        finally:release.set();service.close()
    def test_catalog_counts_and_detail_sources_come_from_saved_data(self):
        service=ThreadsService(self.db,self.config);raw=self.raw()
        try:
            node=node_from_item(item(),relationship='UNRESOLVED');raw.ingest(node,QUERY,'r1');raw.ingest(node,QUERY,'r2')
            catalog=service.catalog();q=next(q for q in catalog['queries'] if q['query_id']==QUERY)
            self.assertEqual(q['saved_nodes'],1);self.assertIsNone(q['last_run'])
            classifier=ClassificationService(self.db)
            workspace=Workspace(self.db,self.db,self.config,classifier,SimpleNamespace(threads=service))
            detail=workspace.detail(node.node_id);self.assertEqual(detail['source_refs'][0]['platform'],'threads')
            self.assertEqual(detail['source_refs'][0]['source_id'],QUERY)
            self.assertIn('threads',workspace.health())
        finally:raw.close();service.close()
    def test_catalog_does_not_report_orphan_run_as_active(self):
        service=ThreadsService(self.db,self.config);raw=self.raw()
        try:
            with raw.transaction():raw.db.execute('INSERT INTO acquisition_runs VALUES (?,?,?,?,?)',('orphan','request','COLLECTING',json.dumps({'run_id':'orphan','status':'COLLECTING','query_snapshot':[]}),utcnow()))
            self.assertIsNone(service.catalog()['current_run'])
        finally:raw.close();service.close()
    def test_run_idempotency_busy_and_cancel(self):
        waiting=threading.Event();released=threading.Event()
        def executor(*args,**kwargs):waiting.set();released.wait(5);return {'status':'CANCELLED' if kwargs['cancel'].is_set() else 'PARTIAL','counters':{}}
        service=ThreadsService(self.db,self.config,executor=executor)
        payload={'project_id':'vinwonders_vu_yen','query_ids':[QUERY],'window_start':'2026-10-01T00:00:00Z','window_end':'2026-10-02T00:00:00Z','idempotency_key':'request'}
        try:
            run=service.start(payload);waiting.wait(2);self.assertEqual(service.start(payload)['run_id'],run['run_id'])
            with self.assertRaises(ThreadsError) as error:service.start({**payload,'idempotency_key':'other'})
            self.assertEqual(error.exception.code,'PLATFORM_BUSY');service.cancel(run['run_id'])
        finally:
            released.set()
            import time
            for _ in range(100):
                if not service.cancels:break
                time.sleep(.02)
            service.close()
    def test_schedule_save_does_not_enable_and_opening_inventory_separate(self):
        service=ThreadsService(self.db,self.config);s=service.schedule_status()
        saved=service.configure_schedule({'expected_version':s['version'],'interval_minutes':45})
        self.assertFalse(saved['enabled'])
        service.configure_schedule({'expected_version':saved['version'],'enabled':True})
        with patch.object(service,'start',return_value={'run_id':'fake'}) as start:
            service.tick(datetime(2026,10,1,22,0,tzinfo=timezone.utc))
            self.assertIsNone(start.call_args.args[1]);self.assertTrue(start.call_args.args[3])
        service.close()
    def test_platform_and_query_scope_filters_and_export(self):
        raw=self.raw()
        try:raw.ingest(node_from_item(item(text='Không nhận được hỗ trợ')),QUERY,'run',lineage={'query_id':QUERY,'project_scope_id':'vinwonders_vu_yen'});raw.ingest(ContentNode('POST',platform_item_id='fb',raw_text='fb'),'fb_group_139108668844789','fb')
        finally:raw.close()
        workspace=Workspace(self.db,self.db,self.config,ClassificationService(self.db),SimpleNamespace())
        view=workspace.view({'kind':'LIBRARY','platform':'threads','project_id':'vinwonders_vu_yen','run_id':'run','date_basis':'published_at'})
        self.assertEqual(view['total'],1);self.assertEqual(view['rows'][0]['platform'],'threads')
        book=load_workbook(BytesIO(workspace.export(view['view_id'])));headers=[c.value for c in book.active[1]]
        self.assertIn('Nền tảng',headers);self.assertIsNone(book.active.cell(2,headers.index('Group')+1).value)
        self.assertEqual(book.active.cell(2,headers.index('Nguồn / truy vấn')+1).value,'VinWonders Vũ Yên')
    def test_api_denied_rate_limit_redirect_never_followed(self):
        for status,code in [(401,'TOKEN_OR_PERMISSION_DENIED'),(429,'HTTP_429'),(302,'THREADS_API_HTTP_ERROR')]:
            response=SimpleNamespace(status_code=status,close=lambda:None)
            class HTTP:
                def get(self,*args,**kwargs):self.options=kwargs;return response
            http=HTTP();adapter=ThreadsAPIAdapter('secret',Deadline(5),http)
            with self.assertRaises(AdapterError) as error:adapter._get('keyword_search',{})
            self.assertEqual(error.exception.code,code);self.assertFalse(http.options['allow_redirects']);self.assertNotIn('secret',str(error.exception))
    def test_parallel_platform_producers_and_review_do_not_overwrite(self):
        self.raw().close();barrier=threading.Barrier(2);errors=[]
        def produce(platform):
            raw=ConversationStore(self.db)
            try:
                barrier.wait(timeout=3)
                for i in range(8):
                    value=node_from_item(item('t'+str(i))) if platform=='THREADS' else ContentNode('POST',platform_item_id=str(i),raw_text='Facebook')
                    result=raw.ingest(value,QUERY if platform=='THREADS' else 'fb','parallel')
                    raw.review(result['node_id'],'RELEVANT','ghi chú '+platform)
            except Exception as error:errors.append(error)
            finally:raw.close()
        workers=[threading.Thread(target=produce,args=(p,)) for p in ('FACEBOOK','THREADS')]
        for worker in workers:worker.start()
        for worker in workers:worker.join(timeout=10)
        self.assertFalse(any(w.is_alive() for w in workers));self.assertFalse(errors)
        raw=self.raw()
        try:
            self.assertEqual(len(raw.list_nodes()),16);self.assertEqual(raw.db.execute('SELECT count(*) FROM node_reviews').fetchone()[0],16)
        finally:raw.close()

if __name__=='__main__':unittest.main()
