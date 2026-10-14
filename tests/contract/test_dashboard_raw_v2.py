"""Raw-first dashboard routing, independent child time, retention and review/export."""
import tempfile
import json
import time
import unittest
from pathlib import Path
from unittest.mock import patch
from io import BytesIO
from openpyxl import load_workbook
from engine.dashboard import SearchService
from engine.dashboard_conversations import crawl, results, pending_run
from engine.conversation_library import library
from engine.library_export import export_library
from engine.registry import Registry
from src.browser.profiles import PROJECT_ROOT
from src.storage.conversations import ConversationStore
from src.domain.conversation import ContentNode
from src.acquisition.models import Window, parse_utc

CONFIG = PROJECT_ROOT / 'config'
SOURCE = 'fb_group_139108668844789'
POST = 'https://www.facebook.com/groups/139108668844789/posts/123'


def ingest(raw):
    root = raw.ingest(ContentNode('POST', platform_item_id='123', canonical_url=POST,
        raw_text='Tuyển dụng công nhân, liên hệ 0901234567',
        published_at='2026-10-01T01:00:00Z',publication_precision='INSTANT'),SOURCE,'fixture')['node_id']
    child = raw.ingest(ContentNode('COMMENT',platform_item_id='456',canonical_url=POST+'?comment_id=456',
        root_post_id=root,parent_node_id=root,raw_text='Tôi chờ 40 phút chưa vào được cổng',
        published_at='2026-10-09T01:00:00Z',publication_precision='INSTANT'),SOURCE,'fixture')['node_id']
    raw.ingest(ContentNode('REPLY',platform_item_id='789',canonical_url=POST+'?comment_id=456&reply_comment_id=789',
        root_post_id=root,parent_node_id=child,raw_text=None),SOURCE,'fixture')
    return child


class DashboardRawTests(unittest.TestCase):
    def test_source_failure_after_checkpoint_preserves_raw_progress_and_results(self):
        import sqlite3
        with tempfile.TemporaryDirectory() as folder:
            db=Path(folder)/'raw.db';registry=Registry(CONFIG)
            window=Window(parse_utc('2026-10-09T00:00:00Z'),parse_utc('2026-10-10T00:00:00Z'))
            def failing(registry,raw,task,**options):
                raw.ingest(ContentNode('POST',platform_item_id='new',canonical_url=POST+'4',
                    raw_text='Kẹt cổng',published_at='2026-10-09T02:00:00Z',publication_precision='INSTANT'),SOURCE,'partial_inventory')
                options['progress']('COMMENTS',dict(source_id=SOURCE,status='PARTIAL',inventory={'run_id':'partial_inventory'},threads=[]))
                raise sqlite3.OperationalError('database is locked')
            with patch('engine.dashboard_conversations.execute_task',failing):
                run=crawl(registry,db,[SOURCE],window,'failed_after_raw',executor=failing)
            self.assertEqual(run['status'],'PARTIAL');self.assertEqual(run['progress']['observed_nodes'],1)
            self.assertEqual(run['jobs'][0]['error_type'],'OperationalError')
            self.assertEqual(len(results(db,window,registry.sources,{},report=run)),1)

    def test_slow_inventory_leaves_time_for_comments(self):
        from engine.conversation_worker import execute_task
        from engine.conversation_schedule import CrawlTasks
        from unittest.mock import Mock
        with tempfile.TemporaryDirectory() as folder:
            raw=ConversationStore(Path(folder)/'raw.db');ingest(raw)
            tasks=CrawlTasks(raw);post=next(n for n in raw.list_nodes() if n['node_type']=='POST')
            tasks.add_post(SOURCE,post['node_id'],post['canonical_url'])
            remaining=[120]
            clock=Mock();clock.remaining.side_effect=lambda:remaining[0]
            def inventory(*args,**options):
                remaining[0]-=args[4] # Simulate consuming the entire listing allowance.
                return {'run_id':'inventory-test','status':'SUCCESS'}
            task=dict(task_id='test',source_id=SOURCE,inventory_only=False,window_start='2026-10-09T00:00:00Z',window_end='2026-10-10T00:00:00Z')
            try:
                with patch('engine.conversation_worker.Deadline',return_value=clock),patch('engine.conversation_worker.inventory_source',side_effect=inventory),patch('engine.conversation_worker.trace_thread',return_value={'status':'PARTIAL'}) as trace:
                    result=execute_task(Registry(CONFIG),raw,task,seconds=120)
                    trace.assert_called_once();self.assertGreater(trace.call_args.args[5],60)
                    self.assertEqual(len(result['threads']),1)
            finally:raw.close()

    def test_run_progress_and_results_use_observations_not_source_history(self):
        from src.storage.conversations import encode
        with tempfile.TemporaryDirectory() as folder:
            db=Path(folder)/'raw.db';registry=Registry(CONFIG)
            raw=ConversationStore(db);ingest(raw);raw.close()
            window=Window(parse_utc('2026-10-09T00:00:00Z'),parse_utc('2026-10-10T00:00:00Z'))
            def executor(registry,raw,task,**options):
                raw.ingest(ContentNode('POST',platform_item_id='new',canonical_url=POST+'4',
                    raw_text='Kẹt cổng',published_at='2026-10-09T02:00:00Z',publication_precision='INSTANT'),SOURCE,'this_inventory')
                return dict(source_id=SOURCE,status='PARTIAL',inventory={'run_id':'this_inventory'})
            snapshots=[]
            run=crawl(registry,db,[SOURCE],window,'this_run',executor,progress_callback=lambda r:snapshots.append(json.loads(encode(r))))
            self.assertEqual([r['current_phase'] for r in snapshots],['INVENTORY','SOURCE_FINISHED','FINISHED'])
            self.assertEqual(run['progress']['completed_sources'],1)
            self.assertEqual(run['progress']['observed_nodes'],1)
            self.assertEqual(run['progress']['confirmed_window_nodes'],1)
            self.assertEqual(run['unresolved_time'],0) # Historical unknown reply belongs to another run.
            rows=results(db,window,registry.sources,{},report=run)
            self.assertEqual(len(rows),1);self.assertEqual(rows[0]['node_type'],'POST')
            self.assertEqual(len(results(db,window,registry.sources,{})),2)

    def test_running_job_publishes_progress_and_read_failure_keeps_it(self):
        import threading
        with tempfile.TemporaryDirectory() as folder:
            service=SearchService(Path(folder)/'raw.db',CONFIG)
            entered=threading.Event();release=threading.Event()
            def blocked(*args,progress_callback,**options):
                report=dict(crawl_run_id=args[4],status='PARTIAL',jobs=[{'source_id':SOURCE,'threads':[{'error_code':'NAVIGATION_TIMEOUT'}]}],unresolved_time=2,
                    updated_at='2026-10-09T01:00:00Z',current_source=SOURCE,current_phase='COMMENTS',
                    progress={'completed_sources':0,'observed_nodes':3,'unknown_time_nodes':2})
                progress_callback(report);entered.set();release.wait(3);return report
            try:
                with patch('engine.dashboard_conversations.crawl',side_effect=blocked),patch('engine.dashboard_conversations.results',side_effect=RuntimeError('read unavailable')):
                    job=service.start({'start_date':'2026-10-09','end_date':'2026-10-09'},'CRAWL')
                    self.assertTrue(entered.wait(2));running=service.get(job['job_id'])
                    self.assertEqual(running['status'],'RUNNING');self.assertEqual(running['progress']['observed_nodes'],3)
                    self.assertEqual(running['warnings'],['NAVIGATION_TIMEOUT'])
                    release.set();deadline=time.monotonic()+3
                    while service.running and time.monotonic()<deadline:time.sleep(.01)
                    failed=service.get(job['job_id']);self.assertEqual(failed['status'],'FAILED')
                    self.assertEqual(failed['progress']['observed_nodes'],3);self.assertIn('CRAWL_FAILED',failed['warnings'])
            finally:release.set();service.close()

    def test_resume_keeps_job_identity_across_more_than_cache_limit_retries(self):
        with tempfile.TemporaryDirectory() as folder:
            service=SearchService(Path(folder)/'raw.db',CONFIG)
            identifier='a'*32
            report=dict(crawl_run_id='dashboard_'+identifier,source_ids=[SOURCE],
                        status='PARTIAL',jobs=[],unresolved_time=0)
            window=Window(parse_utc('2026-10-09T00:00:00Z'),parse_utc('2026-10-10T00:00:00Z'))
            try:
                with patch('engine.dashboard_conversations.crawl',return_value=report),patch('engine.dashboard_conversations.results',return_value=[]):
                    for _ in range(25):
                        job=service.start({},'CRAWL',window_override=window,resume_report=report)
                        self.assertEqual(job['job_id'],identifier)
                        deadline=time.monotonic()+3
                        while service.running and time.monotonic()<deadline:time.sleep(.01)
                        self.assertFalse(service.running)
                    self.assertEqual(len(service.jobs),1)
                    self.assertEqual(service.get(identifier)['crawl_run_id'],report['crawl_run_id'])
            finally:service.close()

    def test_pending_run_reader_does_not_request_writer_lock(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Path(folder)/'raw.db';raw=ConversationStore(db)
            raw.db.execute('CREATE TABLE dashboard_raw_runs(run_id TEXT PRIMARY KEY,payload TEXT NOT NULL)')
            raw.db.commit()
            with raw.transaction():self.assertIsNone(pending_run(db))
            raw.close()

    def test_scheduler_survives_transient_database_lock(self):
        import sqlite3
        from unittest.mock import Mock
        with tempfile.TemporaryDirectory() as folder:
            service=SearchService(Path(folder)/'raw.db',CONFIG,enable_scheduler=False)
            try:
                service.schedule['enabled']=True
                stop=Mock();stop.wait.side_effect=[False,False,True]
                service.scheduler_stop=stop
                error=sqlite3.OperationalError('database is locked');error.sqlite_errorcode=sqlite3.SQLITE_BUSY
                with patch('engine.dashboard_conversations.pending_run',side_effect=[error,None]),patch.object(service,'schedule_tick') as tick:
                    service._schedule_loop()
                    tick.assert_called_once()
                self.assertTrue(service.schedule['enabled'])
            finally:service.close()
    def test_completed_run_is_idempotent_and_busy_profile_remains_resumable(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Path(folder)/'raw.db'; registry=Registry(CONFIG)
            window=Window(parse_utc('2026-10-09T00:00:00Z'),parse_utc('2026-10-10T00:00:00Z'))
            calls=[]
            def executor(registry,raw,task,**kwargs):
                calls.append(task['task_id']); ingest(raw)
                return dict(source_id=SOURCE,status='PARTIAL')
            crawl(registry,db,[SOURCE],window,'done',executor)
            crawl(registry,db,[SOURCE],window,'done',executor)
            self.assertEqual(len(calls),1);self.assertIsNone(pending_run(db))
            def busy(*args,**kwargs): return dict(source_id=SOURCE,status='SKIPPED',error_code='PROFILE_IN_USE')
            crawl(registry,db,[SOURCE],window,'busy',busy)
            self.assertIsNone(pending_run(db))
            with patch('engine.dashboard_conversations.utcnow',return_value='2099-01-01T00:00:00Z'):
                self.assertEqual(pending_run(db)['crawl_run_id'],'busy')

    def test_old_commercial_parent_does_not_hide_new_comment_and_unknown_reply_is_retained(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder); db=base/'raw.db'; raw=ConversationStore(db)
            child=ingest(raw); raw.review(child,'RELEVANT','Needs investigation'); raw.close()
            window=Window(parse_utc('2026-10-09T00:00:00Z'),parse_utc('2026-10-10T00:00:00Z'))
            registry=Registry(CONFIG)
            rows=results(db,window,registry.sources,{})
            self.assertEqual([r['node_type'] for r in rows],['COMMENT'])
            retained=library(base/'legacy.db',db,sources=registry.sources)
            self.assertEqual(retained['total'],3)
            self.assertEqual({r['node_type'] for r in retained['rows']},{'POST','COMMENT','REPLY'})
            record=next(r for r in retained['rows'] if r['node_type']=='COMMENT')
            self.assertEqual(record['review']['note'],'Needs investigation')
            self.assertEqual(record['platform'],'facebook')
            self.assertEqual(library(base/'legacy.db',db,{'node_type':'COMMENT'},registry.sources)['total'],1)
            with self.assertRaises(ValueError): library(base/'legacy.db',db,{'node_type':'INVALID'})
            workbook=load_workbook(BytesIO(export_library(base/'legacy.db',sources=registry.sources,conversation_db=db)))
            self.assertEqual(workbook.active.max_row,4)

    def test_bounded_source_failure_keeps_partial_raw_and_durable_tasks(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Path(folder)/'raw.db'; registry=Registry(CONFIG)
            def executor(registry,raw,task,**kwargs):
                ingest(raw)
                return dict(source_id=task['source_id'],status='PARTIAL',coverage={'status':'UNKNOWN'})
            window=Window(parse_utc('2026-10-09T00:00:00Z'),parse_utc('2026-10-10T00:00:00Z'))
            report=crawl(registry,db,[SOURCE],window,'test_run',executor)
            self.assertFalse(report['full_window_coverage_verified'])
            raw=ConversationStore(db)
            self.assertEqual(raw.db.execute('SELECT status FROM conversation_tasks').fetchone()[0],'DONE')
            self.assertEqual(len(raw.list_nodes()),3); raw.close()

    def test_crawl_and_schedule_dispatch_only_raw_v2_not_acquisition_engine(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder)
            service=SearchService(base/'legacy.db',CONFIG,conversation_db=base/'raw.db')
            report=dict(status='PARTIAL',crawl_run_id='v2',jobs=[],unresolved_time=0)
            with patch('engine.dashboard_conversations.crawl',return_value=report) as worker, \
                 patch('engine.dashboard_conversations.results',return_value=[]), \
                 patch('engine.dashboard.AcquisitionEngine') as legacy:
                job=service.start({'start_date':'2026-10-09','end_date':'2026-10-09'},'CRAWL')
                deadline=time.monotonic()+5
                while service.running and time.monotonic()<deadline: time.sleep(.01)
                self.assertFalse(service.running)
                worker.assert_called_once(); legacy.assert_not_called()
                self.assertEqual(service.get(job['job_id'])['pipeline'],'RAW_V2_THREADS')
                service.configure_schedule(dict(enabled=True,window_mode='DAILY_SLOTS'))
                service.schedule['next_run_at']=None
                service.schedule_tick(parse_utc('2026-10-09T10:45:00Z'))
                while service.running and time.monotonic()<deadline: time.sleep(.01)
                self.assertEqual(worker.call_count,2); legacy.assert_not_called()
            service.close()
