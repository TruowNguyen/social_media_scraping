import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
from engine.classification_service import ClassificationService
from engine.dashboard_workspace import Workspace, WorkspaceError
from engine.worker_heartbeat import Heartbeat
from src.domain.conversation import ContentNode
from src.storage.conversations import ConversationStore

ROOT=Path(__file__).resolve().parents[2]
SOURCE='fb_group_139108668844789'
URL='https://www.facebook.com/groups/139108668844789/posts/123'


def fixture(db,count=3):
    raw=ConversationStore(db)
    root=raw.ingest(ContentNode('POST',platform_item_id='123',canonical_url=URL,raw_text='Bài gốc về cổng dự án',
        published_at='2026-10-01T01:00:00Z',publication_precision='INSTANT'),SOURCE,'seed',observed_at='2026-10-09T01:00:00Z')['node_id']
    ids=[root]
    for i in range(count-1):
        ids.append(raw.ingest(ContentNode('COMMENT',platform_item_id=str(456+i),canonical_url=URL+f'?comment_id={456+i}',
            root_post_id=root,parent_node_id=root,raw_text='Tôi chờ 40 phút ở cổng',author_name='Người dùng',author_uid='001234',
            published_at='2026-10-09T01:00:00Z',publication_precision='INSTANT'),SOURCE,'seed',observed_at='2026-10-09T01:00:00Z')['node_id'])
    raw.close();return ids


class WorkspaceTests(unittest.TestCase):
    def test_ocr_detail_filter_keyword_and_full_excel_evidence(self):
        from src.storage.image_evidence import initialize,save
        raw=ConversationStore(self.db);node=raw.get(self.ids[0]);ref={'media_type':'IMAGE','url':'https://scontent.example.fbcdn.net/photo.jpg'}
        raw.ingest(ContentNode('POST',platform_item_id='123',canonical_url=URL,raw_text=node['raw_text'],media_refs=[ref],published_at=node['published_at'],publication_precision='INSTANT'),SOURCE,'image')
        initialize(raw);text='OCR_SEARCH_TOKEN '+('x'*66000)
        save(raw,self.ids[0],ref,dict(status='AVAILABLE',quality_status='ACCEPTABLE',text=text,processed_at='2026-10-10T01:00:00Z'))
        raw.close()
        view=self.view(ocr_status='AVAILABLE',keyword='OCR_SEARCH_TOKEN')
        self.assertEqual(view['total'],1);self.assertEqual(self.w.detail(self.ids[0])['node']['ocr_text'],text)
        book=load_workbook(BytesIO(self.w.export(view['view_id'])))
        self.assertEqual(''.join(r[6] for r in list(book['Chữ trong ảnh'].values)[1:]),text)
        self.assertEqual(book['Kho dữ liệu']['X2'].data_type,'s')

    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.db=Path(self.tmp.name)/'raw.db';self.ids=fixture(self.db)
        self.w=Workspace(self.db,self.db,ROOT/'config',ClassificationService(self.db),SimpleNamespace())
    def tearDown(self):self.tmp.cleanup()

    def test_health_counts_current_nodes_separately_from_job_history(self):
        raw=ConversationStore(self.db)
        from engine.classification_jobs import ClassificationJobs
        ClassificationJobs(raw)
        raw.close()
        statuses=['DONE','PENDING','MISSING_TEXT']
        with patch('engine.classification_worker.classification_rows',return_value=[{'processing_status':s} for s in statuses]):
            health=self.w.health()
        self.assertEqual(health['classification']['counts'],dict.fromkeys(statuses,1))
        self.assertEqual(health['classification']['job_counts'],{})
        self.assertEqual(sum(health['classification']['counts'].values()),health['nodes'])
    def view(self,**filters):return self.w.view({'kind':'LIBRARY','date_basis':'published_at','refresh':'1',**filters})
    def change(self,id,note='ghi chú'):
        d=self.w.detail(id)
        return dict(node_id=id,revision_id=d['node']['revision_id'],expected_version=d['review']['version'],status='RELEVANT',note=note)
    def test_frozen_membership_review_pagination_and_export_formula(self):
        first=self.view(page_size=1)
        self.w.save_reviews({'request_id':'first','changes':[self.change(self.ids[1],'=HYPERLINK("bad")')]})
        fixture(self.db,5)
        same=self.w.view({'view_id':first['view_id'],'page':2,'page_size':1})
        self.assertEqual(same['total'],3)
        self.assertEqual(same['rows'][0]['review']['status'],'UNREVIEWED')
        fresh=self.view();self.assertEqual(fresh['total'],5)
        wb=load_workbook(BytesIO(self.w.export(fresh['view_id'])))
        cell=next(c for line in wb.active for c in line if c.value=='=HYPERLINK("bad")')
        self.assertEqual(cell.data_type,'s');self.assertEqual(wb.active.max_row,6)
        self.assertIn('Phạm vi',wb.sheetnames)
    def test_atomic_conflict_and_idempotency(self):
        one,two=self.change(self.ids[0]),self.change(self.ids[1])
        body={'request_id':'a','changes':[one]}
        result=self.w.save_reviews(body);self.assertEqual(self.w.save_reviews(body),result)
        with self.assertRaises(WorkspaceError) as error:self.w.save_reviews({'request_id':'b','changes':[two,one]})
        self.assertEqual(error.exception.status,409)
        self.assertEqual(self.w.detail(self.ids[1])['review']['status'],'UNREVIEWED')
        raw=ConversationStore(self.db,initialize=False)
        self.assertEqual(raw.db.execute('SELECT count(*) FROM node_review_history').fetchone()[0],1);raw.close()
    def test_child_time_and_observation_semantics(self):
        child=self.view(start_date='2026-10-09',end_date='2026-10-09')
        self.assertEqual(child['total'],2)
        raw=ConversationStore(self.db,initialize=False)
        node=raw.get(self.ids[1]);raw.ingest(ContentNode('COMMENT',platform_item_id='456',canonical_url=URL+'?comment_id=456',root_post_id=self.ids[0],parent_node_id=self.ids[0],raw_text=node['raw_text'],published_at=node['published_at'],publication_precision='INSTANT'),SOURCE,'later',observed_at='2026-10-10T01:00:00Z');raw.close()
        any_view=self.view(date_basis='observed_any',start_date='2026-10-09',end_date='2026-10-09')
        latest=self.view(date_basis='observed_latest',start_date='2026-10-09',end_date='2026-10-09')
        self.assertEqual(any_view['total'],3);self.assertEqual(latest['total'],2)
    def test_detail_thread_no_provider_and_unknown_worker(self):
        with patch('engine.classification_service.GeminiGateway') as provider:
            d=self.w.detail(self.ids[1]);self.assertEqual(d['ancestors'][0]['node_id'],self.ids[0])
            thread=self.w.thread(self.ids[0],{'limit':1});self.assertEqual(thread['total'],2);self.assertEqual(thread['next_offset'],1)
            self.view();self.assertEqual(self.w.health()['classification']['health'],'UNKNOWN');provider.assert_not_called()
        with Heartbeat(self.db):self.assertEqual(self.w.health()['classification']['health'],'ACTIVE')
        self.assertEqual(self.w.health()['classification']['health'],'STOPPED')

    def test_source_references_and_missing_text_preview(self):
        raw=ConversationStore(self.db)
        empty=raw.ingest(ContentNode('POST',platform_item_id='99999',raw_text=None),SOURCE,'empty')['node_id']
        raw.close()
        for kind in ('CLASSIFY','LIBRARY'):
            view=self.w.view({'kind':kind,'refresh':'1'})
            row=next(r for r in view['rows'] if r['node_id']==self.ids[0])
            self.assertEqual(row['source_refs'][0]['source_id'],SOURCE)
            self.assertIn('facebook.com/groups/',row['source_refs'][0]['url'])
            missing=next(r for r in view['rows'] if r['node_id']==empty)
            self.assertEqual(missing['processing_status'],'MISSING_TEXT')
        self.assertEqual(self.w.preview()['eligible'],len(self.ids))
        view=self.w.view({'kind':'CLASSIFY','processing_status':'MISSING_TEXT','refresh':'1'})
        self.assertEqual(view['total'],1)

    def test_backend_filter_choices_and_invalid_conditions(self):
        options=self.w.catalog()['filter_options']
        self.assertEqual(len(options['intents']),10)
        self.assertIn('MISSING_TEXT',{i['id'] for i in options['processing_statuses']})
        for condition in ({'intent':'NOT_SUPPORTED'},{'processing_status':'GHOST'},
                          {'issue_id':'FAKE_ISSUE'},{'source_id':'missing'},{'made_up_filter':'yes'}):
            with self.subTest(condition=condition),self.assertRaises(WorkspaceError) as error:
                self.w.view({'kind':'CLASSIFY','refresh':'1',**condition})
            self.assertEqual(error.exception.code,'INVALID_FILTER')

    def test_ocr_filter_summary_uses_same_scope_as_rows(self):
        from src.storage.image_evidence import initialize,save
        raw=ConversationStore(self.db);node=raw.get(self.ids[0]);ref={'media_type':'IMAGE','url':'https://scontent.example.fbcdn.net/p.jpg'}
        raw.ingest(ContentNode('POST',platform_item_id='123',canonical_url=URL,raw_text=node['raw_text'],media_refs=[ref]),SOURCE,'image')
        initialize(raw);save(raw,self.ids[0],ref,dict(status='AVAILABLE',quality_status='ACCEPTABLE',text='Chờ vào cổng',processed_at='2026-10-10T01:00:00Z'));raw.close()
        view=self.w.view({'kind':'CLASSIFY','ocr_status':'AVAILABLE','refresh':'1'})
        self.assertEqual(view['total'],1)
        self.assertEqual(sum(view['summary'].values()),1)
        self.assertEqual(sum(view['processing_counts'].values()),1)

    def test_source_project_filter_keeps_pending_and_export_taxonomy_is_not_summary(self):
        from engine.classification_worker import reconcile_saved,classification_cycle
        from src.analysis.taxonomy import load_snapshot
        from tests.contract.test_classification_worker import Provider
        raw=ConversationStore(self.db);snapshot=load_snapshot(self.w.classifier.taxonomy_dir)
        classification_cycle(raw,snapshot,Provider());raw.close()
        view=self.w.view({'kind':'CLASSIFY','project_id':'vincons','project_basis':'SOURCE','refresh':'1'})
        book=load_workbook(BytesIO(self.w.export(view['view_id'])))
        headings=[c.value for c in book.active[1]]
        self.assertIn('Nhóm vấn đề (taxonomy)',headings);self.assertIn('Ý định nội dung',headings)
        issue_name=next(i['name'] for i in snapshot['catalog']['issues'] if i['issue_id']=='ENTRY_QUEUE')
        self.assertEqual(book.active.cell(2,headings.index('Nhóm vấn đề (taxonomy)')+1).value,issue_name)
        raw=ConversationStore(self.db)
        new=raw.ingest(ContentNode('POST',platform_item_id='5555',raw_text='Chờ vào cổng rất lâu'),SOURCE,'pending')['node_id']
        reconcile_saved(raw,snapshot,Provider.model);raw.close()
        view=self.w.view({'kind':'CLASSIFY','project_id':'vincons','project_basis':'SOURCE','refresh':'1'})
        self.assertIn(new,{r['node_id'] for r in view['rows']})
        self.assertIsNone(next(r for r in view['rows'] if r['node_id']==new)['analysis'])
    def test_expired_view_and_retry_guard(self):
        view=self.view();raw=ConversationStore(self.db,initialize=False)
        raw.db.execute("UPDATE dashboard_views SET expires_at='2000-01-01T00:00:00Z'");raw.db.commit();raw.close()
        with self.assertRaises(WorkspaceError) as error:self.w.stored_view(view['view_id'])
        self.assertEqual(error.exception.status,410)
        with self.assertRaises(WorkspaceError):self.w.retry({'node_ids':[self.ids[0]]})
        self.assertEqual(self.w.preview()['scope'],'GLOBAL')

    def test_missing_parent_fallback_and_unknown_publication(self):
        raw=ConversationStore(self.db,initialize=False)
        orphan=raw.ingest(ContentNode('COMMENT',platform_item_id='999',display_url=URL,url_type='PARENT_FALLBACK',
            relationship_status='UNRESOLVED',root_post_id='missing_root',parent_node_id='missing_parent',raw_text='Phản ánh chưa rõ ngày'),SOURCE,'orphan')['node_id'];raw.close()
        detail=self.w.detail(orphan)
        self.assertIn('ROOT_NOT_COLLECTED',detail['missing_context']);self.assertIn('PARENT_NOT_COLLECTED',detail['missing_context'])
        self.assertEqual(detail['link_precision'],'POST_FALLBACK')
        all_rows=self.view();self.assertEqual(all_rows['total'],4)
        dated=self.view(start_date='2026-10-09',end_date='2026-10-09');self.assertEqual(dated['total'],2)


if __name__=='__main__':unittest.main()
