"""Export retains filtered evidence, typed dates and saved review independently of paging."""
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from engine.library_export import export_library, HEADERS
from src.acquisition.models import Candidate
from src.storage.acquisition import AcquisitionStore


class LibraryExportTests(unittest.TestCase):
    def test_all_pages_saved_review_unicode_and_text_safety(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'test.sqlite3'
            store = AcquisitionStore(path)
            for index in range(21):
                url = f'https://www.facebook.com/groups/123/posts/{index}'
                candidate = Candidate('facebook', url, url, 'group', 'run', 'BROWSER',
                                      platform_item_id=str(index), caption='=SUM(1,2)',
                                      published_at='2026-10-08T18:00:00Z', publication_precision='INSTANT',
                                      collected_at='2026-10-08T18:00:00Z')
                rid, _ = store.save_candidate(candidate)
                if index == 0:
                    store.review_content(rid, 'RELEVANT', '=HYPERLINK("https://example.com")\u0001 Ghi ch\u00fa')
            store.close()
            workbook = load_workbook(BytesIO(export_library(path, {'page': '2'})))
            sheet = workbook.active
            self.assertEqual(sheet.max_row, 22)
            self.assertEqual([cell.value for cell in sheet[1]], HEADERS)
            self.assertEqual(sheet['F2'].value.hour, 1)
            self.assertEqual(sheet['F2'].value.day, 9)
            self.assertEqual(sheet['E2'].data_type, 's')
            self.assertEqual(sheet['E2'].value, '=SUM(1,2)')
            self.assertEqual(sheet['L2'].data_type, 's')
            self.assertIn('Ghi ch\u00fa', sheet['L2'].value)
            self.assertNotIn('\u0001', sheet['L2'].value)
            self.assertEqual(sheet.freeze_panes, 'C2')
            self.assertEqual(sheet.auto_filter.ref, 'A1:Q22')
            self.assertFalse(any(c.data_type == 'f' for row in sheet for c in row))
            filtered = load_workbook(BytesIO(export_library(path, {'review': 'RELEVANT', 'start_date': '2026-10-09', 'end_date': '2026-10-09'}))).active
            self.assertEqual(filtered.max_row, 2)
            self.assertEqual(filtered['K2'].value, 'Li\u00ean quan')
            empty = load_workbook(BytesIO(export_library(path, {'keyword': 'absent'}))).active
            self.assertEqual(empty.max_row, 1)
            with self.assertRaises(ValueError): export_library(path, {'start_date': 'invalid'})

    def test_missing_dates_remain_blank(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'test.sqlite3'
            store = AcquisitionStore(path)
            url = 'https://www.facebook.com/groups/123/posts/456'
            store.save_candidate(Candidate('facebook', url, url, 'group', 'run', 'BROWSER'))
            store.close()
            sheet = load_workbook(BytesIO(export_library(path))).active
            self.assertIsNone(sheet['F2'].value)
            self.assertEqual(sheet['G2'].value, 'UNKNOWN')


if __name__ == '__main__': unittest.main()
